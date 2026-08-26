"""Unified report generator.

Key design: each CompareMode defines its own column layout.
The generator consumes unified CompareResult + CompareStats,
eliminating the if mode == ... branching in the original code.
"""

import os
import re
import time

from ..models import CompareResult, CompareStats, CompareMode, CompareStatus


# ---------- Column layout per mode ----------
MODE_COLUMNS = {
    CompareMode.EXCEL_VS_PDF: {
        "sheet_title": "PDF→Excel 器件核对",
        "columns": [
            ("PDF位号", "key"),
            ("核对结果", "status"),
            ("Excel值(Value)", "extra.valueA"),
            ("Excel封装(Footprint)", "extra.footA"),
            ("Excel数量", "extra.qty"),
            ("PDF附近标注", "extra.near"),
        ],
    },
    CompareMode.EXCEL_VS_PDF_MPN: {
        "sheet_title": "MPN 对比(料号)",
        "columns": [
            ("MPN(料号)", "key"),
            ("核对结果", "status"),
            ("Excel数量", "extra.qty_a"),
            ("PDF数量", "extra.qty_b"),
            ("Excel位号", "extra.refs_a"),
            ("PDF位号", "extra.refs_b"),
            ("厂商", "extra.mfg"),
            ("PDF页", "extra.pages"),
        ],
    },
    CompareMode.EXCEL_VS_EXCEL_MPN: {
        "sheet_title": "Excel MPN 对比",
        "columns": [
            ("MPN(物料号)", "key"),
            ("结果", "status"),
            ("数量A", "source_a.qty"),
            ("数量B", "source_b.qty"),
            ("位号A", "source_a.refs"),
            ("位号B", "source_b.refs"),
            ("厂商A", "source_a.mfg"),
            ("厂商B", "source_b.mfg"),
            ("差异明细", "diffs_text"),
        ],
    },
    CompareMode.EXCEL_VS_EXCEL: {
        "sheet_title": "对比明细",
        "columns": [
            ("器件/位号", "key"),
            ("结果", "status"),
            ("文件A行号/页码", "extra.a"),
            ("文件B行号/页码", "extra.b"),
        ],
    },
    CompareMode.PDF_VS_PDF: {
        "sheet_title": "PDF 对比",
        "columns": [
            ("器件/位号", "key"),
            ("结果", "status"),
            ("PDF A 页码", "extra.a_pages"),
            ("PDF B 页码", "extra.b_pages"),
        ],
    },
    CompareMode.GROUP: {
        "sheet_title": "分组对比",
        "columns": [
            ("物料分组", "key"),
            ("结果", "status"),
            ("位号A", "extra.refs_a"),
            ("位号B", "extra.refs_b"),
            ("数量A", "extra.qty_a"),
            ("数量B", "extra.qty_b"),
        ],
    },
}


def _get_nested(obj, path: str, default=""):
    for part in path.split("."):
        if obj is None:
            return default
        if isinstance(obj, dict):
            obj = obj.get(part, default)
        elif hasattr(obj, part):
            obj = getattr(obj, part, default)
        else:
            return default
    return obj if obj is not None else default


def _resolve_value(result: CompareResult, key_path: str) -> str:
    if key_path == "key":
        return result.key
    if key_path == "status":
        return result.status.value
    if key_path == "diffs_text":
        return "；".join(
            f"{d.field}: {d.value_a}→{d.value_b}"
            for d in result.diffs
        )
    if key_path.startswith("source_a."):
        val = _get_nested(result.source_a, key_path[9:])
        if isinstance(val, list):
            return ",".join(str(v) for v in val)
        return str(val) if val else ""
    if key_path.startswith("source_b."):
        val = _get_nested(result.source_b, key_path[9:])
        if isinstance(val, list):
            return ",".join(str(v) for v in val)
        return str(val) if val else ""
    if key_path.startswith("extra."):
        val = _get_nested(result.extra, key_path[6:])
        if isinstance(val, list):
            return ",".join(str(v) for v in val)
        return str(val) if val else ""
    return str(_get_nested(result, key_path))


def _esc_html(s):
    return (
        str(s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def generate_report(
    results: list[CompareResult],
    stats: CompareStats,
    out_dir: str,
    source_label: str = "",
    path_a: str = "",
    path_b: str = "",
) -> tuple[str, str]:
    import openpyxl
    from openpyxl.styles import PatternFill, Font

    os.makedirs(out_dir, exist_ok=True)
    ts = time.strftime("%Y%m%d_%H%M%S")
    mode = stats.mode
    mode_cfg = MODE_COLUMNS.get(mode, MODE_COLUMNS[CompareMode.EXCEL_VS_EXCEL])

    xlsx_path = os.path.join(out_dir, f"对比报告_{ts}.xlsx")
    txt_path = os.path.join(out_dir, f"对比报告_{ts}.txt")

    # ---------- Excel report ----------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = mode_cfg["sheet_title"]

    red_fill = PatternFill("solid", fgColor="FFC7CE")
    red_font = Font(color="9C0006", bold=True)

    headers = [h for h, _ in mode_cfg["columns"]]
    ws.append(headers)

    for r in results:
        row_vals = [_resolve_value(r, kp) for _, kp in mode_cfg["columns"]]
        ws.append(row_vals)
        if not r.is_match:
            r_ = ws.max_row
            for cc in range(1, len(headers) + 1):
                c_ = ws.cell(row=r_, column=cc)
                c_.fill = red_fill
                c_.font = red_font

    # Summary sheet
    ws2 = wb.create_sheet("汇总")
    ws2.append(["统计项", "数量"])
    for k, v in stats.summary.items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["仅A有：", stats.only_a])
    for r in results:
        if r.is_only_a:
            ws2.append(["", r.key])
    ws2.append(["仅B有：", stats.only_b])
    for r in results:
        if r.is_only_b:
            ws2.append(["", r.key])

    wb.save(xlsx_path)

    # ---------- TXT report ----------
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("=" * 70 + "\n")
        f.write(f"文件对比报告 ({mode.value})\n")
        f.write(f"来源: {source_label}\n")
        f.write(f"时间: {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write("=" * 70 + "\n\n")
        f.write("【汇总】\n")
        for k, v in stats.summary.items():
            f.write(f"  {k:<30s}: {v}\n")

        f.write("\n【明细】\n")
        col_widths = [max(len(h), 20) for h in headers]
        header_line = " | ".join(h.ljust(w) for h, w in zip(headers, col_widths))
        f.write(header_line + "\n")
        f.write("-" * len(header_line) + "\n")
        for r in results:
            vals = [_resolve_value(r, kp) for _, kp in mode_cfg["columns"]]
            line = " | ".join(v.ljust(w)[:w] for v, w in zip(vals, col_widths))
            f.write(line + "\n")

        f.write("\n【仅A有】\n")
        for r in results:
            if r.is_only_a:
                f.write(f"  {r.key}\n")
        f.write("\n【仅B有】\n")
        for r in results:
            if r.is_only_b:
                f.write(f"  {r.key}\n")

    _generate_html(txt_path, results, stats, mode_cfg["columns"])
    _generate_pdf(out_dir, txt_path, results, stats, path_a, path_b)

    return xlsx_path, txt_path


def _generate_html(txt_path, results, stats, columns):
    html_path = os.path.splitext(txt_path)[0] + ".html"
    try:
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(
                '<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8">'
                "<title>BOM 核对报告</title>"
                "<style>body{font-family:'Microsoft YaHei',sans-serif;margin:24px;color:#23303B}"
                "h1{color:#1E4E84}table{border-collapse:collapse;width:100%}"
                "td,th{border:1px solid #dfe6f0;padding:6px 10px;font-size:13px;text-align:left}"
                "th{background:#E8F0FA;color:#1E4E84}.ok{background:#E7F7EA;color:#1E7B34}"
                ".bad{background:#FDEAEA;color:#C0392B}.warn{background:#FFF4E0;color:#996A00}</style>"
                "</head><body>"
            )
            f.write("<h1>BOM 核对 / 对比报告</h1>")
            f.write(f"<p>生成时间: {time.strftime('%Y-%m-%d %H:%M:%S')}</p>")
            f.write("<h2>汇总</h2><table><tr><th>统计项</th><th>数量</th></tr>")
            for k, v in stats.summary.items():
                f.write(f"<tr><td>{_esc_html(k)}</td><td>{_esc_html(v)}</td></tr>")
            f.write("</table>")

            f.write("<h2>明细</h2><table><tr>")
            for h, _ in columns:
                f.write(f"<th>{_esc_html(h)}</th>")
            f.write("</tr>")
            for r in results:
                cls = "ok" if r.is_match else "bad"
                f.write(f"<tr class='{cls}'>")
                for _, kp in columns:
                    f.write(f"<td>{_esc_html(_resolve_value(r, kp))}</td>")
                f.write("</tr>")
            f.write("</table></body></html>")
    except Exception:
        pass


# ---------- PDF report (Arena-style, matches original exactly) ----------
def _clip(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[:n] + "…"


def _generate_pdf(out_dir, txt_path, results, stats, path_a, path_b):
    """Generate Arena-style PDF report matching 835&685对比.pdf format.
    Only for Excel-Excel MPN mode.
    """
    if stats.mode != CompareMode.EXCEL_VS_EXCEL_MPN:
        return

    try:
        import pymupdf
    except ImportError:
        return

    base = os.path.splitext(os.path.basename(txt_path))[0]
    pdf_path = os.path.join(out_dir, base + ".pdf")
    page_w, page_h = 1191.0, 842.0
    RED = (0.6, 0, 0)
    BLACK = (0, 0, 0)
    BODY = "helv"
    BOLD = "hebo"
    ITAL = "hebi"

    doc = pymupdf.open()
    page = doc.new_page(width=page_w, height=page_h)

    def _p(x, y, text, size=8, font=BODY, color=BLACK):
        page.insert_text((x, y), text, fontsize=size, fontname=font, color=color)

    # ---- 顶部页眉 ----
    _p(34, 55, "Seyond Inc.", 11.2, BOLD)
    _p(34, 68, "Seyond Production", 9, BOLD)
    _p(34, 80, "Printed by", 6, BOLD)
    _p(34, 90, f"Printed on {time.strftime('%m/%d/%Y')}", 6, BOLD)
    _p(34, 100, "Local time zone (GMT+08:00) China Standard Time", 6, BOLD)
    _p(34, 112, "Page  >  Bill of Materials  >  Compare", 8)
    _p(34, 121, "All information contained in this document is proprietary and confidential.",
       6, ITAL, (0, 0.067, 0.09))

    # ---- 文件标题 ----
    name_a = os.path.basename(path_a or "")
    name_b = os.path.basename(path_b or "")
    _p(50, 145, name_a, 13.5, BOLD)
    _p(50, 163, "Rev A1", 8)
    _p(680, 145, name_b, 13.5, BOLD)
    _p(680, 163, "Rev A1", 8)

    # ---- 表头 ----
    y0 = 251.0
    _p(35, y0, ">>", 8, BODY, RED)
    _p(48, y0, "#", 8, BOLD)
    _p(115, y0, "ITEM NUMBER", 8, BOLD)
    _p(284, y0, "ITEM NAME", 8, BOLD)
    _p(647, y0 - 4, "PCS", 7.5, BOLD)
    _p(743, y0, "REF DES", 8, BOLD)
    _p(1103, y0, "SUBSTITUTES", 8, BOLD)

    page.draw_line(
        pymupdf.Point(30, y0 + 6), pymupdf.Point(1165, y0 + 6),
        color=(0.8, 0.8, 0.8), width=0.6,
    )

    # ---- 数据行 ----
    row_h = 15.0
    y = y0 + 16
    n = 0
    for r in results:
        if y > page_h - 40:
            page = doc.new_page(width=page_w, height=page_h)
            y = 40
        n += 1
        changed = not r.is_match
        ra = r.source_a or {}
        c = RED if changed else BLACK
        if changed:
            _p(35, y, ">>", 8, BOLD, RED)
        _p(48, y, f"{n:02d}", 8, BODY, c)

        itm_a = str(ra.get("itm") or "") or str(r.key)
        _p(115, y, _clip(itm_a, 34), 8, BODY, c)
        _p(284, y, _clip(str(ra.get("desc") or ""), 70), 8, BODY, c)

        qa = str(ra.get("qty") or "")
        _p(647, y, f"PCS{qa}" if qa else "PCS", 8, BODY, c)

        refs = ra.get("refs") or []
        _p(743, y, _clip(",".join(refs), 55), 8, BODY, c)

        if r.status == CompareStatus.A_ONLY:
            sub = "仅文件A(文件B无此料号)"
        elif r.status == CompareStatus.B_ONLY:
            sub = "仅文件B(文件A无此料号)"
        elif r.diffs:
            sub = "；".join(
                f"{d.field}: {d.value_a}→{d.value_b}"
                for d in r.diffs
            )[:60]
        else:
            sub = ""
        _p(1103, y, _clip(sub, 45), 8, BODY, RED if sub else BLACK)
        y += row_h

    doc.save(pdf_path)
    doc.close()