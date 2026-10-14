"""
BRD-AI: DFM 报告生成器 (Valor 式, 内嵌版图截图)
=============================================================
将分析结果输出为:
  - Excel DFM 报告 (*.xlsx):
      Main           封面 (单板信息 / 问题汇总 / 严重度图例 / 风险 / 版图概览图 /
                     层映射 / EDA工具 / 板及装配统计)
      Issue summary  问题清单 (序号/问题点/问题图片/描述/建议/数量/优先级/状态/层/
                     网络/X/Y/潜在损失/损失小计/备注) + 冻结窗格 + 自动筛选 + 分级配色
      Info           制造及装配信息 (工艺统计 + 制程能力备注)
      开短路网表      开短路测试 / 网表一致性 (统计 + 明细含XY)
      Issue1..N      单项问题明细 (规则/层/网络/测量/规则值/X/Y/严重度 + 定位截图 + 全部定位点)
  - 文本 DFM 报告 (*.txt)

内嵌截图: 由 src.dfm_snap (Pillow) 渲染整板概览与逐问题定位放大图; 无几何数据或
未安装 Pillow 时自动跳过, 不影响报告生成。
"""
import math
import os
import shutil
import tempfile
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.dfm_rules import SEVERITY_LEVELS, SEVERITY_COLOR

FONT = "Microsoft YaHei"

# 严重度 -> 潜在损失估算 (相对值, 仅供排序/参考)
LOSS_BY_SEV = {"危险": 2000, "高": 800, "中": 300, "低": 50, "关闭": 0}
SEVERITY_DESC = {
    "危险": "停产 / 可靠性影响",
    "高": "成本增加 / 可靠性影响",
    "中": "影响成本或制造良率",
    "低": "建议优化",
    "关闭": "曾违规项已关闭",
}
EDA_TOOLS = ["Cadence Allegro", "Mentor Xpedition", "Mentor Pads", "Altium Designer",
             "Zuken CR-8000", "Valor NPI", "Valor VUV", "Valor DFA"]


def _thin_border():
    return Border(left=Side(style="thin", color="BFBFBF"),
                  right=Side(style="thin", color="BFBFBF"),
                  top=Side(style="thin", color="BFBFBF"),
                  bottom=Side(style="thin", color="BFBFBF"))


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4g}"
    return str(value)


def _numfmt(value: Any) -> Any:
    """X/Y/测量值等数值保留原始数字类型以便 Excel 排序; 非数值原样返回"""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 4)
    return value


def _write_kv(ws, start_row, col, pairs, key_font=None, border=None, val_fmt=None):
    r = start_row
    for key, val in pairs:
        kc = ws.cell(row=r, column=col, value=key)
        kc.font = key_font or Font(name=FONT, bold=True, size=10)
        vc = ws.cell(row=r, column=col + 1, value=val_fmt(val) if val_fmt else _fmt(val))
        vc.font = Font(name=FONT, size=10)
        if border:
            kc.border = border
            vc.border = border
        r += 1
    return r


# ============================================================
# 截图准备
# ============================================================
def _prepare_images(analysis: dict, tmpdir: str) -> dict:
    imgs = {"overview": None, "crops": {}}
    fab = analysis.get("fab_data")
    if not fab:
        return imgs
    try:
        from src.dfm_snap import render_overview_png, render_crop_png, _collect_bbox
    except Exception:  # noqa: BLE001  (无 Pillow 等)
        return imgs
    issues = analysis.get("issues", [])
    try:
        imgs["overview"] = render_overview_png(
            fab, [it for it in analysis.get("instances", issues) if it.get("x") is not None],
            os.path.join(tmpdir, "overview.png"), width=560)
    except Exception:  # noqa: BLE001
        imgs["overview"] = None
    bbox = _collect_bbox(fab)
    if bbox:
        diag = math.hypot(bbox[2] - bbox[0], bbox[3] - bbox[1])
        half = max(3.0, min(20.0, diag * 0.05))
    else:
        half = 6.0
    for idx, it in enumerate(issues, 1):
        if it.get("x") is None or it.get("y") is None:
            continue
        try:
            imgs["crops"][idx] = render_crop_png(
                fab, it["x"], it["y"], os.path.join(tmpdir, f"crop_{idx}.png"),
                half_span_mm=half, size=320)
        except Exception:  # noqa: BLE001
            pass
    return imgs


def _embed(ws, path, anchor, px):
    if not path or not os.path.isfile(path):
        return False
    try:
        from openpyxl.drawing.image import Image as XLImage
        xi = XLImage(path)
        xi.width = px
        xi.height = px
        ws.add_image(xi, anchor)
        return True
    except Exception:  # noqa: BLE001
        return False


# ============================================================
# Excel 报告
# ============================================================
def generate_dfm_excel(analysis: dict, output_path: str,
                       board_meta: Optional[dict] = None,
                       embed_images: bool = True) -> str:
    board_meta = board_meta or {}
    issues: List[dict] = analysis.get("issues", [])
    risk = analysis.get("risk", {})
    info_stats = analysis.get("info_stats", {})
    summary = analysis.get("issue_summary", {})

    wb = openpyxl.Workbook()
    border = _thin_border()
    title_font = Font(name=FONT, size=16, bold=True, color="1F3864")
    sub_font = Font(name=FONT, size=10, color="404040")
    header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    section_font = Font(name=FONT, size=11, bold=True, color="1F3864")
    base_font = Font(name=FONT, size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    top_wrap = Alignment(vertical="top", wrap_text=True)

    def style_header_row(ws, row, c1, c2):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

    tmpdir = tempfile.mkdtemp(prefix="dfm_img_") if embed_images else None
    imgs = _prepare_images(analysis, tmpdir) if tmpdir else {"overview": None, "crops": {}}

    # ==================== Main (封面) ====================
    ws = wb.active
    ws.title = "Main"
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 3), ("B", 18), ("C", 34), ("D", 3), ("E", 16),
                   ("F", 16), ("G", 3), ("H", 14)):
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:H1")
    ws["B1"] = "DFM 分析报告  (Design For Manufacturability)"
    ws["B1"].font = title_font
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells("B2:H2")
    ws["B2"] = (f"板厂: {analysis.get('factory_label','')}    "
                f"板类型: {analysis.get('board_type_label','')}    "
                f"能力等级: {analysis.get('capability_label','')}    "
                f"铜厚: {_fmt(analysis.get('copper_um'))}um")
    ws["B2"].font = sub_font
    ws.merge_cells("B3:H3")
    ws["B3"] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}    工具: Asher_DFM (参考 Valor)"
    ws["B3"].font = sub_font

    # 单板信息 (左)
    ws["B5"] = "单板信息"
    ws["B5"].font = section_font
    board_pairs = [
        ("客户名", board_meta.get("customer", "")),
        ("单板名", board_meta.get("board_name", "")),
        ("单板版本", board_meta.get("version", "")),
        ("装配名", board_meta.get("assembly_name", "")),
        ("规则集", board_meta.get("rule_set", "")),
        ("EDA同步工具", board_meta.get("eda_tool", "Cadence Allegro")),
        ("提交时间", board_meta.get("submit_time", "")),
        ("完成时间", board_meta.get("finish_time", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))),
        ("提交者", board_meta.get("submitter", "Asher_DFM")),
        ("问题总数", summary.get("总计", len(issues))),
        ("已解决总数", sum(1 for it in issues if it.get("status") == "关闭")),
    ]
    _write_kv(ws, 6, 2, board_pairs, border=border)

    # 问题汇总 (右上)
    ws["E5"] = "问题汇总"
    ws["E5"].font = section_font
    r = 6
    for s in SEVERITY_LEVELS:
        c1 = ws.cell(row=r, column=5, value=s)
        c1.fill = PatternFill("solid", fgColor=SEVERITY_COLOR.get(s, "FFFFFF"))
        c1.font = Font(name=FONT, size=10, bold=True)
        c1.border = border
        c1.alignment = center
        c2 = ws.cell(row=r, column=6, value=summary.get(s, 0))
        c2.border = border
        c2.font = base_font
        c2.alignment = center
        r += 1
    c1 = ws.cell(row=r, column=5, value="总计")
    c1.font = Font(name=FONT, size=10, bold=True)
    c1.border = border
    c1.alignment = center
    c2 = ws.cell(row=r, column=6, value=summary.get("总计", len(issues)))
    c2.font = Font(name=FONT, size=10, bold=True)
    c2.border = border
    c2.alignment = center

    # 严重度图例 (右, 汇总下方)
    leg_row = r + 2
    ws.cell(row=leg_row, column=5, value="严重度图例").font = section_font
    rr = leg_row + 1
    for s in SEVERITY_LEVELS:
        c1 = ws.cell(row=rr, column=5, value=s)
        c1.fill = PatternFill("solid", fgColor=SEVERITY_COLOR.get(s, "FFFFFF"))
        c1.font = Font(name=FONT, size=9, bold=True)
        c1.alignment = center
        c1.border = border
        ws.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=8)
        c2 = ws.cell(row=rr, column=6, value=SEVERITY_DESC.get(s, ""))
        c2.font = Font(name=FONT, size=9)
        c2.alignment = left_wrap
        rr += 1

    # 风险 (左, 单板信息下方)
    risk_row = 18
    ws.cell(row=risk_row, column=2, value="可制造性评估").font = section_font
    risk_pairs = [
        ("可制造性评分", f"{risk.get('score','-')} / 100"),
        ("风险等级", risk.get('risk_level', '-')),
        ("预估良率区间", f"{risk.get('yield_estimate',['-','-'])[0]}% ~ "
                          f"{risk.get('yield_estimate',['-','-'])[1]}%"),
        ("综合结论", risk.get('verdict', '-')),
    ]
    _write_kv(ws, risk_row + 1, 2, risk_pairs, border=border)

    # 版图概览图 (左, 风险下方)
    ov_row = risk_row + 6
    ws.cell(row=ov_row, column=2, value="版图概览 (问题定位)").font = section_font
    if imgs.get("overview"):
        _embed(ws, imgs["overview"], f"B{ov_row + 1}", 330)
        img_bottom = ov_row + 1 + int(330 / 15) + 1
    else:
        ws.cell(row=ov_row + 1, column=2,
                value="(无几何数据, 未生成概览图)").font = sub_font
        img_bottom = ov_row + 3

    # 层映射 (右, 与概览图并排)
    lm_row = ov_row
    ws.cell(row=lm_row, column=5, value="层映射 / 叠层").font = section_font
    stack = analysis.get("layer_stack") or []
    copper = analysis.get("copper_layers") or []
    rr = lm_row + 1
    if stack:
        for c, h in ((5, "层"), (6, "类型"), (7, "厚度mm")):
            ws.cell(row=rr, column=c, value=h)
        style_header_row(ws, rr, 5, 7)
        rr += 1
        for st in stack:
            ws.cell(row=rr, column=5, value=st.get("layer", "")).border = border
            ws.cell(row=rr, column=6, value=st.get("type", "")).border = border
            ws.cell(row=rr, column=7, value=_fmt(st.get("thickness_mm"))).border = border
            for c in (5, 6, 7):
                ws.cell(row=rr, column=c).font = base_font
            rr += 1
    elif copper:
        for nm in copper:
            ws.cell(row=rr, column=5, value=nm).font = base_font
            rr += 1

    # EDA 工具 (右, 层映射下方)
    rr += 1
    ws.cell(row=rr, column=5, value="支持的 EDA / 同步工具").font = section_font
    rr += 1
    for i in range(0, len(EDA_TOOLS), 2):
        ws.cell(row=rr, column=5, value=EDA_TOOLS[i]).font = base_font
        if i + 1 < len(EDA_TOOLS):
            ws.cell(row=rr, column=6, value=EDA_TOOLS[i + 1]).font = base_font
        rr += 1

    # 板及装配统计 (底部, 概览图下方)
    stat_row = max(img_bottom, rr) + 2
    ws.cell(row=stat_row, column=2, value="单板及装配统计信息").font = section_font
    man = info_stats.get("manufacturing", {})
    asm = info_stats.get("assembly", {})
    _write_kv(ws, stat_row + 1, 2, list(man.items()), border=border)
    _write_kv(ws, stat_row + 1, 5, list(asm.items()), border=border)

    # ==================== Issue summary ====================
    ws2 = wb.create_sheet("Issue summary")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:O1")
    ws2["A1"] = "DFM 问题清单"
    ws2["A1"].font = title_font
    ws2.row_dimensions[1].height = 24
    ws2.merge_cells("A2:O2")
    ws2["A2"] = (f"问题总数: {summary.get('总计', len(issues))}    "
                 f"危险 {summary.get('危险',0)} / 高 {summary.get('高',0)} / "
                 f"中 {summary.get('中',0)} / 低 {summary.get('低',0)} / 关闭 {summary.get('关闭',0)}    "
                 f"(潜在损失为按严重度的相对估算)")
    ws2["A2"].font = sub_font

    headers = ["序号", "问题点(code)", "问题图片", "问题描述", "专家建议", "数量",
               "优先级", "状态", "涉及层", "网络", "X位置mm", "Y位置mm",
               "潜在损失(估)", "损失小计(估)", "备注"]
    widths = [6, 24, 13, 44, 44, 7, 9, 8, 12, 14, 11, 11, 13, 13, 16]
    start = 4
    for c, h in enumerate(headers, 1):
        ws2.cell(row=start, column=c, value=h)
        ws2.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    style_header_row(ws2, start, 1, len(headers))
    ws2.row_dimensions[start].height = 28

    total_loss = 0
    for idx, it in enumerate(issues, 1):
        r = start + idx
        sev = it.get("severity", "")
        cnt = it.get("count", 1) or 1
        loss = LOSS_BY_SEV.get(sev, 0)
        subtotal = loss * cnt
        total_loss += subtotal
        detail_sheet = f"Issue{idx}"

        ws2.cell(row=r, column=1, value=idx)
        code_cell = ws2.cell(row=r, column=2, value=f"{it.get('rule','')} [{it.get('code','')}]")
        code_cell.hyperlink = f"#'{detail_sheet}'!A1"
        code_cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        # 问题图片 (缩略图)
        crop = imgs.get("crops", {}).get(idx)
        if crop:
            _embed(ws2, crop, f"C{r}", 84)
            ws2.row_dimensions[r].height = 66
        else:
            ws2.cell(row=r, column=3, value="—").alignment = center
            ws2.row_dimensions[r].height = 30
        ws2.cell(row=r, column=4, value=it.get("description", ""))
        ws2.cell(row=r, column=5, value=it.get("suggestion", ""))
        ws2.cell(row=r, column=6, value=cnt)
        sev_cell = ws2.cell(row=r, column=7, value=sev)
        sev_cell.fill = PatternFill("solid", fgColor=SEVERITY_COLOR.get(sev, "FFFFFF"))
        sev_cell.font = Font(name=FONT, size=10, bold=True)
        sev_cell.alignment = center
        ws2.cell(row=r, column=8, value=it.get("status", ""))
        ws2.cell(row=r, column=9, value=it.get("layer", ""))
        ws2.cell(row=r, column=10, value=it.get("net", ""))
        xc = ws2.cell(row=r, column=11, value=_numfmt(it.get("x")))
        yc = ws2.cell(row=r, column=12, value=_numfmt(it.get("y")))
        xc.number_format = "0.000"
        yc.number_format = "0.000"
        ws2.cell(row=r, column=13, value=loss)
        ws2.cell(row=r, column=14, value=subtotal)
        ws2.cell(row=r, column=15, value=it.get("note", ""))
        for c in range(1, len(headers) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.border = border
            if c not in (2, 7):          # 2=超链接字体, 7=严重度字体, 已单独设置
                cell.font = base_font
            if c == 7:
                continue
            if c in (4, 5, 15):
                cell.alignment = top_wrap
            elif c in (1, 3, 6, 8, 11, 12, 13, 14):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    last_row = start + len(issues)
    if issues:
        tr = last_row + 1
        ws2.cell(row=tr, column=4, value="潜在损失合计(估)").font = Font(name=FONT, bold=True, size=10)
        tc = ws2.cell(row=tr, column=14, value=total_loss)
        tc.font = Font(name=FONT, bold=True, size=10)
        tc.border = border
    ws2.freeze_panes = "A5"
    if issues:
        ws2.auto_filter.ref = f"A{start}:O{last_row}"

    # ==================== Info ====================
    ws3 = wb.create_sheet("Info")
    ws3.sheet_view.showGridLines = False
    for col, w in (("A", 3), ("B", 22), ("C", 16), ("D", 10), ("E", 3),
                   ("F", 22), ("G", 16)):
        ws3.column_dimensions[col].width = w
    ws3.merge_cells("B1:G1")
    ws3["B1"] = "DFM 报告 - 制造及装配信息"
    ws3["B1"].font = title_font
    ws3["B3"] = "制造工艺"
    ws3["B3"].font = section_font
    for c, h in ((2, "规格"), (3, "数值"), (4, "单位")):
        ws3.cell(row=4, column=c, value=h)
    style_header_row(ws3, 4, 2, 4)
    units = {"最小孔径": "mm", "单板厚度": "mm", "微孔最小孔径": "mm",
             "外层线宽": "mm", "外层线距": "mm", "内层线宽": "mm", "内层线距": "mm",
             "单板Pin密度": "Pin/cm²", "拼板铆孔密度": "Via/cm²", "单板尺寸": "mm",
             "板厚孔径比": ":1"}
    r = 5
    for k, v in info_stats.get("manufacturing", {}).items():
        ws3.cell(row=r, column=2, value=k).border = border
        ws3.cell(row=r, column=3, value=_fmt(v)).border = border
        ws3.cell(row=r, column=4, value=units.get(k, "")).border = border
        for c in (2, 3, 4):
            ws3.cell(row=r, column=c).font = base_font
        r += 1
    ws3["F3"] = "装配工艺"
    ws3["F3"].font = section_font
    ws3.cell(row=4, column=6, value="详细资料")
    ws3.cell(row=4, column=7, value="数值")
    style_header_row(ws3, 4, 6, 7)
    r2 = 5
    for k, v in info_stats.get("assembly", {}).items():
        ws3.cell(row=r2, column=6, value=k).border = border
        ws3.cell(row=r2, column=7, value=_fmt(v)).border = border
        for c in (6, 7):
            ws3.cell(row=r2, column=c).font = base_font
        r2 += 1
    notes = info_stats.get("capability_notes", {})
    if notes:
        nr = max(r, r2) + 2
        ws3.cell(row=nr, column=2, value="制程能力备注").font = section_font
        _write_kv(ws3, nr + 1, 2, list(notes.items()), border=border)

    # ==================== 开短路 / 网表测试 ====================
    nl = analysis.get("netlist_summary")
    if nl:
        wsn = wb.create_sheet("开短路网表")
        wsn.sheet_view.showGridLines = False
        wsn.merge_cells("A1:H1")
        wsn["A1"] = "DFM 报告 - 开短路测试 / 网表一致性"
        wsn["A1"].font = title_font
        for col, w in (("A", 10), ("B", 20), ("C", 18), ("D", 12),
                       ("E", 12), ("F", 12), ("G", 14), ("H", 56)):
            wsn.column_dimensions[col].width = w
        wsn["A3"] = "网表统计"
        wsn["A3"].font = section_font
        stat_pairs = [
            ("逻辑网络数", nl.get("logical_net_count")),
            ("物理连通域数", nl.get("physical_component_count")),
            ("开路网络数(Open)", nl.get("open_net_count")),
            ("短路网络对(Short)", nl.get("short_net_count")),
            ("无铜节点数", nl.get("nodes_without_copper")),
            ("几何元素数", nl.get("element_count")),
        ]
        fab_cmp = nl.get("fab_compare")
        if fab_cmp:
            stat_pairs += [("CAD网络数", fab_cmp.get("cad_net_count")),
                           ("FAB网络数", fab_cmp.get("fab_net_count")),
                           ("匹配网络数", fab_cmp.get("matched_nets")),
                           ("CAD-FAB开路", fab_cmp.get("open_count")),
                           ("CAD-FAB短路", fab_cmp.get("short_count"))]
        rr = _write_kv(wsn, 4, 1, stat_pairs, border=border)
        passed = not nl.get("open_net_count") and not nl.get("short_net_count")
        vc = wsn.cell(row=rr + 1, column=1, value="结论")
        vc.font = Font(name=FONT, bold=True, size=10)
        verdict = "PASS - 未发现开路/短路" if passed else "FAIL - 存在开路/短路, 详见下表"
        vcell = wsn.cell(row=rr + 1, column=2, value=verdict)
        vcell.font = Font(name=FONT, bold=True, size=10,
                          color="0A7A3A" if passed else "C00000")

        inst = [i for i in analysis.get("instances", [])
                if i.get("rule") in ("开路(Open)", "短路(Short)", "网络缺失", "多余网络")]
        hdr_row = rr + 3
        wsn.cell(row=hdr_row, column=1, value="开短路 / 网表明细 (含XY定位)").font = section_font
        cols = ["类型", "网络", "从…", "到…", "X位置mm", "Y位置mm", "涉及层", "描述"]
        for c, h in enumerate(cols, 1):
            wsn.cell(row=hdr_row + 1, column=c, value=h)
        style_header_row(wsn, hdr_row + 1, 1, len(cols))
        r0 = hdr_row + 2
        for k, it in enumerate(inst):
            row = r0 + k
            vals = [it.get("rule", ""), it.get("net", ""), it.get("from", ""),
                    it.get("to", ""), _numfmt(it.get("x")), _numfmt(it.get("y")),
                    it.get("layer", ""), it.get("description", "")]
            for c, v in enumerate(vals, 1):
                cell = wsn.cell(row=row, column=c, value=v)
                cell.border = border
                cell.font = base_font
                cell.alignment = top_wrap if c == 8 else Alignment(vertical="center")
                if c in (5, 6):
                    cell.number_format = "0.000"
        if not inst:
            gc = wsn.cell(row=r0, column=1, value="(无开路/短路)")
            gc.font = base_font
            gc.border = border
        wsn.freeze_panes = f"A{r0}"

    # ==================== Issue 明细页 ====================
    detail_cols = ["序号", "规则名称", "涉及层", "网络", "从…开始", "到…结束",
                   "测量值mm", "规则值mm", "X位置mm", "Y位置mm",
                   "严重度", "状态", "首次版本", "客户/研发反馈"]
    for idx, it in enumerate(issues, 1):
        wsd = wb.create_sheet(f"Issue{idx}")
        wsd.sheet_view.showGridLines = False
        wsd.merge_cells("A1:N1")
        wsd["A1"] = f"DFM 报告 - 问题 {idx}  [{it.get('code','')}]"
        wsd["A1"].font = title_font
        wsd.column_dimensions["A"].width = 6
        wsd.column_dimensions["B"].width = 26
        for i in range(3, len(detail_cols) + 1):
            wsd.column_dimensions[get_column_letter(i)].width = 14

        wsd["A3"] = "问题描述:"
        wsd["A3"].font = Font(name=FONT, bold=True, size=10)
        wsd.merge_cells("B3:N3")
        wsd["B3"] = it.get("description", "")
        wsd["B3"].font = base_font
        wsd["B3"].alignment = left_wrap
        wsd["A4"] = "专家建议:"
        wsd["A4"].font = Font(name=FONT, bold=True, size=10)
        wsd.merge_cells("B4:N4")
        wsd["B4"] = it.get("suggestion", "")
        wsd["B4"].font = base_font
        wsd["B4"].alignment = left_wrap

        for c, h in enumerate(detail_cols, 1):
            wsd.cell(row=6, column=c, value=h)
        style_header_row(wsd, 6, 1, len(detail_cols))
        row = 7
        sev = it.get("severity", "")
        vals = [1, f"{it.get('rule','')} [{it.get('code','')}]", it.get("layer", ""),
                it.get("net", ""), it.get("from", ""), it.get("to", ""),
                _numfmt(it.get("measured")), _numfmt(it.get("required")),
                _numfmt(it.get("x")), _numfmt(it.get("y")), sev,
                it.get("status", ""), board_meta.get("version", ""), it.get("note", "")]
        for c, v in enumerate(vals, 1):
            cell = wsd.cell(row=row, column=c, value=v)
            cell.border = border
            cell.font = base_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if c in (7, 8, 9, 10):
                cell.number_format = "0.0000" if c in (7, 8) else "0.000"
        sev_cell = wsd.cell(row=row, column=11)
        sev_cell.fill = PatternFill("solid", fgColor=SEVERITY_COLOR.get(sev, "FFFFFF"))
        sev_cell.font = Font(name=FONT, size=10, bold=True)

        # 定位截图
        crop = imgs.get("crops", {}).get(idx)
        snap_row = row + 2
        wsd.cell(row=snap_row, column=1, value="定位截图:").font = Font(name=FONT, bold=True, size=10)
        if crop:
            _embed(wsd, crop, f"B{snap_row}", 300)
            snap_bottom = snap_row + int(300 / 15) + 2
        else:
            wsd.cell(row=snap_row, column=2, value="(无坐标/无几何, 未生成截图)").font = sub_font
            snap_bottom = snap_row + 2

        # 全部定位点
        locs = it.get("locations") or []
        if len(locs) > 1:
            lr = snap_bottom + 1
            wsd.cell(row=lr, column=1,
                     value=f"全部定位点 (共 {it.get('count', len(locs))} 处, 列出前 {min(len(locs),300)}):"
                     ).font = Font(name=FONT, bold=True, size=10)
            for c, h in ((1, "#"), (2, "X位置mm"), (3, "Y位置mm")):
                wsd.cell(row=lr + 1, column=c, value=h)
            style_header_row(wsd, lr + 1, 1, 3)
            for k, (lx, ly) in enumerate(locs[:300], start=2):
                rr2 = lr + k
                wsd.cell(row=rr2, column=1, value=k - 1).border = border
                xc = wsd.cell(row=rr2, column=2, value=_numfmt(lx))
                yc = wsd.cell(row=rr2, column=3, value=_numfmt(ly))
                xc.number_format = "0.000"
                yc.number_format = "0.000"
                xc.border = border
                yc.border = border
                for c in (1, 2, 3):
                    wsd.cell(row=rr2, column=c).font = base_font

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    if tmpdir:
        shutil.rmtree(tmpdir, ignore_errors=True)
    return output_path


# ============================================================
# 文本报告
# ============================================================
def generate_dfm_text(analysis: dict, output_path: str,
                      board_meta: Optional[dict] = None) -> str:
    board_meta = board_meta or {}
    issues = analysis.get("issues", [])
    risk = analysis.get("risk", {})
    lines = []
    lines.append("=" * 70)
    lines.append("  DFM 制程能力分析报告   (Asher_DFM, 参考 Valor)")
    lines.append("=" * 70)
    lines.append(f"  板厂: {analysis.get('factory_label','')}    "
                 f"板类型: {analysis.get('board_type_label','')}    "
                 f"能力等级: {analysis.get('capability_label','')}")
    lines.append(f"  板名: {board_meta.get('board_name', analysis.get('board_type_label',''))}    "
                 f"铜厚: {_fmt(analysis.get('copper_um'))}um")
    lines.append(f"  生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("-" * 70)
    summary = analysis.get("issue_summary", {})
    lines.append(f"  问题总数: {summary.get('总计', len(issues))}   "
                 f"危险 {summary.get('危险',0)} / 高 {summary.get('高',0)} / "
                 f"中 {summary.get('中',0)} / 低 {summary.get('低',0)}")
    lines.append(f"  可制造性评分: {risk.get('score','-')}/100    "
                 f"风险等级: {risk.get('risk_level','-')}    "
                 f"预估良率: {risk.get('yield_estimate',['-','-'])[0]}%~"
                 f"{risk.get('yield_estimate',['-','-'])[1]}%")
    lines.append(f"  结论: {risk.get('verdict','-')}")
    lines.append("=" * 70)
    if issues:
        lines.append("")
        lines.append("  问题清单:")
        for idx, it in enumerate(issues, 1):
            lines.append(f"  [{it.get('severity','')}] {idx:>2}. {it.get('code','')} "
                         f"{it.get('rule','')}")
            lines.append(f"        {it.get('description','')}")
            if it.get("measured") is not None or it.get("required") is not None:
                lines.append(f"        测量值={_fmt(it.get('measured'))}  规则值={_fmt(it.get('required'))}")
            if it.get("x") is not None and it.get("y") is not None:
                loc = f"        位置: X={_fmt(it.get('x'))}mm  Y={_fmt(it.get('y'))}mm"
                if it.get("layer"):
                    loc += f"  层={it.get('layer')}"
                if it.get("net"):
                    loc += f"  网络={it.get('net')}"
                lines.append(loc)
            lines.append(f"        建议: {it.get('suggestion','')}")
    else:
        lines.append("")
        lines.append("  未发现 DFM 问题。")

    nl = analysis.get("netlist_summary")
    if nl:
        lines.append("")
        lines.append("-" * 70)
        lines.append("  开短路测试 / 网表一致性:")
        lines.append(f"    逻辑网络数        {_fmt(nl.get('logical_net_count'))}")
        lines.append(f"    物理连通域数      {_fmt(nl.get('physical_component_count'))}")
        lines.append(f"    开路网络数(Open)  {_fmt(nl.get('open_net_count'))}")
        lines.append(f"    短路网络对(Short) {_fmt(nl.get('short_net_count'))}")
        lines.append(f"    无铜节点数        {_fmt(nl.get('nodes_without_copper'))}")
        fab_cmp = nl.get("fab_compare")
        if fab_cmp:
            lines.append(f"    CAD vs FAB 网表:  CAD {fab_cmp.get('cad_net_count')} / "
                         f"FAB {fab_cmp.get('fab_net_count')} 网络, "
                         f"Open {fab_cmp.get('open_count')} / Short {fab_cmp.get('short_count')} "
                         f"({'PASS' if fab_cmp.get('passed') else 'FAIL'})")
        verdict = "PASS - 未发现开路/短路" if (
            not nl.get("open_net_count") and not nl.get("short_net_count")) else \
            "FAIL - 存在开路/短路, 详见问题清单 NET-* 项 (含XY定位)"
        lines.append(f"    结论: {verdict}")
        inst = [i for i in analysis.get("instances", [])
                if i.get("rule") in ("开路(Open)", "短路(Short)")]
        if inst:
            lines.append("    开短路定位 (前 50 处):")
            for it in inst[:50]:
                xy = (f"({_fmt(it.get('x'))},{_fmt(it.get('y'))})"
                      if it.get("x") is not None else "(-)")
                lines.append(f"      [{it.get('rule')}] {xy} {it.get('description','')}")

    man = analysis.get("info_stats", {}).get("manufacturing", {})
    asm = analysis.get("info_stats", {}).get("assembly", {})
    if man or asm:
        lines.append("")
        lines.append("-" * 70)
        lines.append("  制造工艺 / 装配工艺统计:")
        for k, v in man.items():
            lines.append(f"    {k:16s} {_fmt(v)}")
        for k, v in asm.items():
            lines.append(f"    {k:16s} {_fmt(v)}")
    lines.append("")
    lines.append("=" * 70)

    content = "\n".join(lines)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    return output_path
